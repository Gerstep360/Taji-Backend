from html import escape
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export_xlsx(title, columns, rows, tenant, description):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reporte"
    sheet.append([column.label for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for row_number, row in enumerate(rows, 2):
        sheet.append([ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value for value in row])
        for index in range(1, len(columns) + 1):
            cell = sheet.cell(row=row_number, column=index)
            if isinstance(cell.value, str):
                # Texto de usuarios, nunca fórmulas, enlaces ni macros ejecutables.
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(45, max(18, len(column.label) + 4))
    info = workbook.create_sheet("Configuración")
    for row in [("Reporte", title), ("Condominio", tenant.name),
                ("Generado", timezone.localtime().strftime("%d/%m/%Y %H:%M")),
                ("Criterios", description)]:
        info.append(row)
    for row in info:
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                cell.data_type = "s"
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 80
    output = BytesIO()
    workbook.save(output)
    return HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def export_html(title, columns, rows, tenant, description):
    headings = "".join(f"<th scope='col'>{escape(column.label)}</th>" for column in columns)
    body = "".join("<tr>" + "".join(f"<td>{escape(str(value))}</td>" for value in row) + "</tr>" for row in rows)
    html = f"""<!doctype html><html lang="es"><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{escape(title)}</title><style>
body{{font-family:Arial,sans-serif;color:#162b4d;margin:32px}}h1{{color:#2563eb}}
table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #dbe3ee;padding:10px;text-align:left;white-space:pre-wrap}}
th{{background:#eff6ff}}tr:nth-child(even){{background:#f8fafc}}p{{white-space:pre-wrap}}
</style><h1>{escape(title)}</h1><p>Condominio: {escape(tenant.name)}</p>
<p>{escape(description)}</p><table><thead><tr>{headings}</tr></thead><tbody>{body}</tbody></table></html>"""
    response = HttpResponse(html, content_type="text/html; charset=utf-8")
    response["Content-Security-Policy"] = "default-src 'none'; style-src 'unsafe-inline'; sandbox"
    return response
