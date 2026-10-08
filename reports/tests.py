from datetime import datetime, timedelta, timezone as utc
from io import BytesIO
from unittest.mock import patch

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import AccessToken

from accounts.models import Person, Role, User
from condominiums.models import Condominium, Resident, Sector, Staff, Unit
from security.models import SecurityShift, ShiftHandover, ShiftLogEntry
from tenancy.context import TenantContext
from tenancy.models import TenantMembership


class ReportsTests(APITestCase):
    def setUp(self):
        TenantContext.clear()
        self.addCleanup(TenantContext.clear)
        self.a = Condominium.objects.create(name="Condominio A", slug="report-a")
        self.b = Condominium.objects.create(name="Condominio B", slug="report-b")
        self.admin_role, _ = Role.objects.get_or_create(slug="administrador", defaults={"name": "Administrador"})
        self.guard_role, _ = Role.objects.get_or_create(slug="seguridad", defaults={"name": "Seguridad"})
        self.admin = User.objects.create_user(email="reports@example.test", password="Test-only-123!", role=self.admin_role)
        self.membership = TenantMembership.objects.create(user=self.admin, condominium=self.a, role=self.admin_role, is_default=True)
        self.client.force_authenticate(self.admin)
        self.guards = []
        self.shifts = []
        self.logs = []
        self.handovers = []
        for index, condo in enumerate((self.a, self.b)):
            guard = Staff.objects.create(condominium=condo, staff_type="SECURITY", status="ACTIVE",
                person=Person.objects.create(first_name="Guardia", last_name=condo.name, document_number=f"00{index + 1}"))
            self.guards.append(guard)
            Resident.objects.create(condominium=condo,
                person=Person.objects.create(first_name="Residente", last_name=condo.name))
            sector = Sector.objects.create(condominium=condo, code=f"SEC-{index}", name=f"Sector {condo.name}")
            Unit.objects.create(sector=sector, code=f"UNIT-{index}")
            start = datetime(2026, 10, 8, 3, 30, tzinfo=utc.utc)  # 7 octubre, 23:30 en Bolivia.
            outgoing = SecurityShift.objects.create(condominium=condo, guard_staff=guard,
                scheduled_start=start, scheduled_end=start + timedelta(hours=1))
            incoming = SecurityShift.objects.create(condominium=condo, guard_staff=guard,
                scheduled_start=start + timedelta(hours=1), scheduled_end=start + timedelta(hours=2))
            self.shifts.append(outgoing)
            self.logs.append(ShiftLogEntry.objects.create(shift=outgoing, title=f"Nota {condo.name}",
                description=f"Descripción {condo.name}", occurred_at=start))
            self.handovers.append(ShiftHandover.objects.create(outgoing_shift=outgoing, incoming_shift=incoming,
                summary=f"Resumen {condo.name}"))

    def payload(self, **extra):
        return {"source": "staff", "columns": ["name", "document", "area"], **extra}

    def preview(self, **extra):
        return self.client.post(reverse("reports-preview"), self.payload(**extra), format="json")

    def export(self, **extra):
        return self.client.post(reverse("reports-export"), self.payload(**{"format": "xlsx", **extra}), format="json")

    def test_catalog_scoped_and_does_not_expose_orm_fields(self):
        response = self.client.get(reverse("reports-catalog"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["condominium"]["id"], self.a.pk)
        self.assertEqual(len(response.data["sources"]), 6)
        self.assertEqual(response.data["formats"], ["xlsx", "html"])
        self.assertNotIn("path", response.data["sources"][0]["columns"][0])
        self.assertEqual(response["Cache-Control"], "no-store")

    def test_columns_and_position_are_respected(self):
        response = self.preview(columns=["area", "document", "name"])
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual([item["key"] for item in response.data["columns"]], ["area", "document", "name"])
        self.assertEqual(response.data["rows"], [["Seguridad", "001", "Guardia Condominio A"]])

    def test_every_source_excludes_other_tenant(self):
        for source in ("staff", "residents", "units", "shifts", "logs", "handovers"):
            with self.subTest(source=source):
                response = self.preview(source=source, columns=["id"])
                self.assertEqual(response.status_code, 200, response.data)
                expected = 2 if source == "shifts" else 1
                self.assertEqual(response.data["pagination"]["total"], expected)

    def test_mismatched_legacy_relationships_are_excluded(self):
        broken = SecurityShift.objects.create(condominium=self.a, guard_staff=self.guards[1],
            scheduled_start=timezone.now(), scheduled_end=timezone.now() + timedelta(hours=1))
        ShiftLogEntry.objects.create(shift=broken, description="No debe aparecer")
        ShiftHandover.objects.create(outgoing_shift=broken, incoming_shift=self.shifts[0], summary="No debe aparecer")
        self.handovers[0].incoming_shift = self.shifts[1]
        self.handovers[0].save()
        self.assertEqual(self.preview(source="shifts", columns=["id"]).data["pagination"]["total"], 2)
        self.assertEqual(self.preview(source="logs", columns=["id"]).data["pagination"]["total"], 1)
        self.assertEqual(self.preview(source="handovers", columns=["id"]).data["pagination"]["total"], 0)

    def test_foreign_sector_of_legacy_log_is_not_disclosed(self):
        self.logs[0].sector = Sector.objects.get(condominium=self.b)
        self.logs[0].save()
        self.assertEqual(self.preview(source="logs", columns=["sector"]).data["rows"], [])

    def test_filters_are_combined_with_and(self):
        response = self.preview(filters=[{"field": "name", "operator": "contains", "value": "Guardia"},
                                        {"field": "area", "operator": "eq", "value": "CLEANING"}])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["rows"], [])
        response = self.preview(filters=[{"field": "document", "operator": "ne", "value": "002"}])
        self.assertEqual(len(response.data["rows"]), 1)

    def test_invalid_columns_operators_choices_and_duplicate_sort_rejected(self):
        for data in ({"columns": []}, {"columns": ["name", "name"]}, {"columns": ["person__user__password"]},
                     {"filters": [{"field": "condominium_id", "operator": "eq", "value": str(self.b.pk)}]},
                     {"filters": [{"field": "name", "operator": "regex", "value": ".*"}]},
                     {"filters": [{"field": "area", "operator": "eq", "value": "INVALID"}]},
                     {"ordering": [{"field": "name", "direction": "asc"}, {"field": "name", "direction": "desc"}]}):
            with self.subTest(data=data):
                self.assertEqual(self.preview(**data).status_code, 400)

    def test_filter_dates_use_tenant_timezone_and_include_entire_day(self):
        response = self.preview(source="shifts", columns=["scheduled_start"],
            filters=[{"field": "scheduled_start", "operator": "eq", "value": "2026-10-07"}])
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["rows"], [["07/10/2026 23:30"]])
        response = self.preview(source="shifts", columns=["id"],
            filters=[{"field": "scheduled_start", "operator": "lte", "value": "2026-10-07"}])
        self.assertEqual(response.data["pagination"]["total"], 1)

    def test_invalid_date_and_number_rejected_without_server_error(self):
        for field, value in (("scheduled_start", "2026-02-30"), ("id", "12.5"), ("id", str(2**90))):
            response = self.preview(source="shifts", columns=["id"],
                filters=[{"field": field, "operator": "eq", "value": value}])
            self.assertEqual(response.status_code, 400, response.data)

    def test_empty_values_and_date_field(self):
        self.assertEqual(self.preview(filters=[{"field": "hire_date", "operator": "is_empty"}]).data["pagination"]["total"], 1)
        self.assertEqual(self.preview(filters=[{"field": "phone", "operator": "not_empty"}]).data["pagination"]["total"], 0)
        self.guards[0].hire_date = "2026-10-07"
        self.guards[0].save()
        response = self.preview(columns=["hire_date"], filters=[{"field": "hire_date", "operator": "gte", "value": "2026-10-01"}])
        self.assertEqual(response.data["rows"], [["07/10/2026"]])

    def test_sort_and_pagination(self):
        Staff.objects.create(condominium=self.a, staff_type="CLEANING",
            person=Person.objects.create(first_name="Zeta", last_name="Local"))
        response = self.preview(columns=["name"], ordering=[{"field": "name", "direction": "desc"}], page_size=1)
        self.assertEqual(response.data["rows"], [["Zeta Local"]])
        self.assertEqual(response.data["pagination"]["pages"], 2)
        self.assertEqual(self.preview(columns=["name"], page_size=1, page=2).data["pagination"]["page"], 2)
        self.assertEqual(self.preview(page_size=101).status_code, 400)
        self.assertEqual(self.preview(page=0).status_code, 400)

    def test_guard_resident_and_orphan_denied_on_every_endpoint(self):
        for slug in ("seguridad", "residente", "administrador"):
            role, _ = Role.objects.get_or_create(slug=slug, defaults={"name": slug})
            actor = User.objects.create_user(email=f"{slug}@example.test", password="Test-only-123!", role=role)
            if slug != "administrador":
                TenantMembership.objects.create(user=actor, condominium=self.a, role=role)
            self.client.force_authenticate(actor)
            self.assertEqual(self.client.get(reverse("reports-catalog")).status_code, 403)
            self.assertEqual(self.preview().status_code, 403)
            self.assertEqual(self.export().status_code, 403)

    def test_membership_role_not_global_role_authorizes(self):
        self.membership.role = self.guard_role
        self.membership.save()
        self.assertEqual(self.preview().status_code, 403)

    def test_inactive_membership_or_condominium_denied(self):
        self.membership.is_active = False
        self.membership.save()
        self.assertEqual(self.preview().status_code, 403)
        self.membership.is_active = True
        self.membership.save()
        self.a.is_active = False
        self.a.save()
        self.assertEqual(self.preview().status_code, 403)

    def test_real_bearer_session_and_forged_tenant_denied(self):
        self.client.force_authenticate(user=None)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {AccessToken.for_user(self.admin)}")
        self.assertEqual(self.preview().status_code, 200)
        response = self.client.post(reverse("reports-preview"), self.payload(), format="json", HTTP_X_TENANT_ID=str(self.b.pk))
        self.assertEqual(response.status_code, 403)

    def test_unauthenticated_denied(self):
        self.client.force_authenticate(user=None)
        self.assertEqual(self.preview().status_code, 401)

    def test_excel_preserves_headers_text_document_and_filters(self):
        response = self.export(columns=["document", "name"], filters=[{"field": "area", "operator": "eq", "value": "SECURITY"}])
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(list(workbook["Reporte"].values), [("Documento", "Nombre completo"), ("001", "Guardia Condominio A")])
        self.assertEqual(workbook["Reporte"].freeze_panes, "A2")
        self.assertIn("Condominio A", str(list(workbook["Configuración"].values)))
        self.assertIn("attachment", response["Content-Disposition"])

    def test_excel_never_interprets_user_text_as_formula(self):
        self.guards[0].notes = '=HYPERLINK("https://example.test","Click")'
        self.guards[0].save()
        response = self.export(columns=["notes"], title="=1+1")
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook["Reporte"]["A2"].data_type, "s")
        self.assertEqual(workbook["Configuración"]["B1"].data_type, "s")
        self.assertEqual(workbook["Reporte"]["A2"].value, self.guards[0].notes)

    def test_html_escapes_all_user_content_and_is_scoped(self):
        self.guards[0].notes = '<script>alert("x")</script>'
        self.guards[0].save()
        response = self.client.post(reverse("reports-export"), self.payload(format="html", columns=["notes"], title="<img>"), format="json")
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)
        self.assertIn("&lt;img&gt;", html)
        self.assertNotIn("Condominio B", html)
        self.assertIn("sandbox", response["Content-Security-Policy"])

    def test_export_contains_all_filtered_rows_not_current_page(self):
        Staff.objects.create(condominium=self.a, staff_type="SECURITY",
            person=Person.objects.create(first_name="Segundo", last_name="Local"))
        response = self.export(columns=["name"], page_size=1, page=2)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook["Reporte"].max_row, 3)

    def test_export_limit_requires_more_filters_instead_of_truncating(self):
        with patch("reports.views.EXPORT_LIMIT", 1):
            self.assertFalse(self.preview(source="shifts", columns=["id"]).data["can_export"])
            response = self.export(source="shifts", columns=["id"])
            self.assertEqual(response.status_code, 400)

    def test_both_formats_support_empty_results(self):
        for format in ("xlsx", "html"):
            response = self.client.post(reverse("reports-export"), self.payload(format=format,
                filters=[{"field": "name", "operator": "contains", "value": "inexistente"}]), format="json")
            self.assertEqual(response.status_code, 200)

    def test_export_requires_known_format_and_filter_limits(self):
        self.assertEqual(self.client.post(reverse("reports-export"), self.payload(), format="json").status_code, 400)
        self.assertEqual(self.export(format="csv").status_code, 400)
        self.assertEqual(self.preview(filters=[{"field": "name", "operator": "eq", "value": "A"}] * 11).status_code, 400)
